from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from scripts.build_deliverable_templates import build_all, load_catalog, validate_catalog
from scripts.validate_skill_suite import EXPECTED_SKILLS


ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "templates" / "catalog.json"


REQUIRED_SECTIONS = {
    "managing-data-projects": {"business_outcome", "scope", "domain_route", "dependencies", "decision_rights", "risks_controls", "acceptance_evidence", "next_gate"},
    "handling-data-ethically": {"purpose", "affected_parties", "harms", "recourse", "stop_conditions"},
    "establishing-data-governance": {"charter", "decision_rights", "policy", "exceptions", "scorecard"},
    "designing-data-architecture": {"current_state", "target_state", "transition", "tradeoffs", "conformance"},
    "modeling-data": {"use_cases", "model_level", "grain_keys", "traceability", "review"},
    "operating-data-storage": {"service_scope", "slo", "recovery", "runbook", "go_no_go"},
    "securing-data": {"classification", "threats", "access_design", "control_tests", "residual_risk"},
    "integrating-data": {"interface_contract", "mapping", "replay", "schema_evolution", "reconciliation"},
    "managing-documents-and-content": {"content_scope", "taxonomy", "retention", "legal_hold", "disposition"},
    "managing-reference-and-master-data": {"authority", "global_identity", "survivorship", "merge_split", "distribution"},
    "delivering-data-warehousing-and-bi": {"decisions", "metric_contract", "fact_grain", "history", "business_acceptance"},
    "managing-metadata": {"metadata_product", "glossary", "lineage", "freshness", "impact_analysis"},
    "improving-data-quality": {"purpose", "rules", "baseline", "root_cause", "prevention", "scorecard"},
    "delivering-data-science": {"hypothesis", "data_features", "evaluation", "model_card", "monitoring", "rollback"},
    "assessing-data-management-maturity": {"scope", "rubric", "evidence", "current_target", "roadmap"},
    "organizing-data-management": {"model_options", "selected_model", "role_charter", "raci", "pilot"},
    "leading-data-change": {"case_for_change", "stakeholders", "communications", "adoption", "reinforcement"},
}

REQUIRED_COLUMNS = {
    "managing-data-projects": {"domain", "goal", "dependency", "deliverable", "acceptance_evidence"},
    "handling-data-ethically": {"affected_party", "harm", "control"},
    "establishing-data-governance": {"decision", "accountable", "escalation"},
    "designing-data-architecture": {"capability", "dependency", "owner"},
    "modeling-data": {"entity", "attribute", "definition", "key"},
    "operating-data-storage": {"rto", "rpo", "test_result"},
    "securing-data": {"subject", "data", "action", "decision"},
    "integrating-data": {"source_field", "target_field", "transformation"},
    "managing-documents-and-content": {"content_type", "retention_rule", "disposition"},
    "managing-reference-and-master-data": {"source_id", "global_id", "match_rule"},
    "delivering-data-warehousing-and-bi": {"metric", "definition", "source", "test"},
    "managing-metadata": {"term", "owner", "source", "consumer"},
    "improving-data-quality": {"rule", "denominator", "threshold", "owner"},
    "delivering-data-science": {"feature", "source", "validation", "drift_threshold"},
    "assessing-data-management-maturity": {"criterion", "evidence_strength", "current", "target"},
    "organizing-data-management": {"activity", "accountable", "responsible"},
    "leading-data-change": {"stakeholder", "action", "adoption_metric"},
}


class TemplateCatalogTests(unittest.TestCase):
    def test_all_skills_and_languages_are_present(self) -> None:
        catalog = load_catalog(CATALOG)
        self.assertEqual(set(catalog["skills"]), set(EXPECTED_SKILLS))
        self.assertEqual(validate_catalog(catalog), [])

    def test_missing_translation_is_rejected(self) -> None:
        catalog = deepcopy(load_catalog(CATALOG))
        del catalog["skills"]["managing-data-projects"]["sections"][0]["instruction"]["en"]
        self.assertIn("missing en guidance", " ".join(validate_catalog(catalog)))

    def test_duplicate_field_id_is_rejected(self) -> None:
        catalog = deepcopy(load_catalog(CATALOG))
        catalog["common_fields"].append(deepcopy(catalog["common_fields"][0]))
        self.assertIn("duplicate id", " ".join(validate_catalog(catalog)))

    def test_each_domain_has_its_required_sections(self) -> None:
        catalog = load_catalog(CATALOG)
        for slug, required in REQUIRED_SECTIONS.items():
            with self.subTest(skill=slug):
                actual = {section["id"] for section in catalog["skills"][slug]["sections"]}
                self.assertTrue(required.issubset(actual), required - actual)
                columns = {column["id"] for sheet in catalog["skills"][slug]["sheets"] for column in sheet["columns"]}
                self.assertTrue(REQUIRED_COLUMNS[slug].issubset(columns), REQUIRED_COLUMNS[slug] - columns)

    def test_builds_six_portable_files_per_skill(self) -> None:
        from docx import Document
        from openpyxl import load_workbook

        with TemporaryDirectory() as temporary:
            skills_root = Path(temporary)
            catalog = load_catalog(CATALOG)
            for slug in EXPECTED_SKILLS:
                (skills_root / slug).mkdir()
            outputs = build_all(CATALOG, skills_root)
            self.assertEqual(len(outputs), 102)
            for slug in EXPECTED_SKILLS:
                with self.subTest(skill=slug):
                    directory = skills_root / slug / "templates"
                    spec = catalog["skills"][slug]
                    self.assertEqual(len(list(directory.iterdir())), 6)
                    for lang in ("zh", "en"):
                        markdown = (directory / f"delivery-pack.{lang}.md").read_text(encoding="utf-8")
                        self.assertIn(slug, markdown)
                        self.assertIn(f"report.{lang}.docx", markdown)
                        self.assertIn(f"register.{lang}.xlsx", markdown)
                        document = Document(directory / f"report.{lang}.docx")
                        self.assertGreater(len(document.paragraphs), 5)
                        report_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                        for section in spec["sections"]:
                            self.assertIn(section["id"], report_text)
                        workbook = load_workbook(directory / f"register.{lang}.xlsx")
                        self.assertEqual(len(workbook.worksheets), len(spec["sheets"]) + 1)
                        self.assertFalse(workbook._external_links)
                        for sheet, sheet_spec in zip(workbook.worksheets[1:], spec["sheets"]):
                            self.assertEqual(sheet.title, sheet_spec["name"][lang])
                            self.assertTrue(sheet.freeze_panes)
                            self.assertTrue(sheet.tables or sheet.auto_filter.ref)
                            headers = [cell.value for cell in sheet[1]]
                            self.assertEqual(len(headers), len(sheet_spec["columns"]))
                            for header, column in zip(headers, sheet_spec["columns"]):
                                self.assertIn(column["id"], header)


if __name__ == "__main__":
    unittest.main()
