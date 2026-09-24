#!/usr/bin/env python3
"""Validate and build the bilingual deliverable-template catalog."""

from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import tempfile
from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
from xml.etree import ElementTree

from scripts.validate_skill_suite import EXPECTED_SKILLS


COMMON_FIELD_IDS = {
    "artifact_id", "project", "skill", "version", "date", "status",
    "author", "owner", "approver", "purpose", "business_outcome",
    "scope", "source_inputs", "assumptions", "decisions", "rationale",
    "risks", "exceptions", "open_items", "due_date",
    "acceptance_criteria", "evidence_location", "reviewer_decision", "next_gate",
}


def _missing_languages(value: dict, location: str, errors: list[str], suffix: str = "") -> None:
    for lang in ("zh", "en"):
        if not isinstance(value, dict) or not isinstance(value.get(lang), str) or not value[lang].strip():
            errors.append(f"{location}: missing {lang}{suffix}")


def _duplicate_ids(rows: list[dict]) -> bool:
    ids = [row.get("id") for row in rows]
    return len(ids) != len(set(ids))


def validate_catalog(catalog: dict) -> list[str]:
    errors: list[str] = []
    if catalog.get("schema_version") != 1:
        errors.append("schema_version must be 1")

    fields = catalog.get("common_fields", [])
    if not isinstance(fields, list) or {field.get("id") for field in fields} != COMMON_FIELD_IDS:
        errors.append("common_fields: wrong field ids")
    elif _duplicate_ids(fields):
        errors.append("common_fields: duplicate id")
    for field in fields if isinstance(fields, list) else []:
        _missing_languages(field, f"common_fields.{field.get('id')}", errors)

    skills = catalog.get("skills", {})
    if not isinstance(skills, dict) or set(skills) != set(EXPECTED_SKILLS):
        errors.append("skill set does not match EXPECTED_SKILLS")
    for slug, spec in skills.items() if isinstance(skills, dict) else []:
        for key in ("title", "report_title", "register_title"):
            _missing_languages(spec.get(key), f"{slug}.{key}", errors)
        sections = spec.get("sections", [])
        if not isinstance(sections, list) or not sections:
            errors.append(f"{slug}.sections: empty")
            sections = []
        if _duplicate_ids(sections):
            errors.append(f"{slug}.sections: duplicate id")
        for section in sections:
            location = f"{slug}.sections.{section.get('id')}"
            if not section.get("id"):
                errors.append(f"{location}: missing id")
            _missing_languages(section, location, errors)
            _missing_languages(section.get("instruction"), location, errors, " guidance")

        sheets = spec.get("sheets", [])
        if not isinstance(sheets, list) or not sheets:
            errors.append(f"{slug}.sheets: empty")
            sheets = []
        if _duplicate_ids(sheets):
            errors.append(f"{slug}.sheets: duplicate id")
        for sheet in sheets:
            location = f"{slug}.sheets.{sheet.get('id')}"
            if not sheet.get("id"):
                errors.append(f"{location}: missing id")
            _missing_languages(sheet.get("name"), location, errors)
            _missing_languages(sheet.get("purpose"), location, errors, " guidance")
            for lang in ("zh", "en"):
                name = sheet.get("name", {}).get(lang, "")
                if len(name) > 31 or any(character in name for character in "[]:*?/\\"):
                    errors.append(f"{location}: invalid {lang} Excel sheet name")
            columns = sheet.get("columns", [])
            if not isinstance(columns, list) or not columns:
                errors.append(f"{location}: empty columns")
                columns = []
            if _duplicate_ids(columns):
                errors.append(f"{location}: duplicate column id")
            for column in columns:
                column_location = f"{location}.{column.get('id')}"
                if not column.get("id"):
                    errors.append(f"{column_location}: missing id")
                _missing_languages(column, column_location, errors)
                _missing_languages(column.get("instruction"), column_location, errors, " guidance")
        for lang in ("zh", "en"):
            names = [sheet.get("name", {}).get(lang) for sheet in sheets]
            if len(names) != len(set(names)):
                errors.append(f"{slug}.sheets: duplicate {lang} sheet name")
    return errors


def load_catalog(path: Path) -> dict:
    source = json.loads(path.read_text(encoding="utf-8"))
    try:
        fields = [
            {"id": field_id, "zh": labels[0], "en": labels[1]}
            for field_id, labels in source["common_fields"].items()
        ]
        catalog = {"schema_version": source["schema_version"], "common_fields": fields, "skills": {}}
        for slug, raw_spec in source["skills"].items():
            spec = {}
            for key in ("title", "report_title", "register_title"):
                spec[key] = {"zh": raw_spec[key][0], "en": raw_spec[key][1]}
            spec["sections"] = [
                {
                    "id": section[0], "zh": section[1], "en": section[2],
                    "instruction": {"zh": section[3], "en": section[4]},
                }
                for section in raw_spec["sections"]
            ]
            spec["sheets"] = []
            for sheet in raw_spec["sheets"]:
                columns = []
                for field_id in sheet[5]:
                    labels = source["column_fields"][field_id]
                    columns.append({
                        "id": field_id, "zh": labels[0], "en": labels[1],
                        "instruction": {"zh": labels[2], "en": labels[3]},
                    })
                spec["sheets"].append({
                    "id": sheet[0],
                    "name": {"zh": sheet[1], "en": sheet[2]},
                    "purpose": {"zh": sheet[3], "en": sheet[4]},
                    "columns": columns,
                })
            catalog["skills"][slug] = spec
    except (KeyError, IndexError, TypeError) as error:
        raise ValueError(f"invalid catalog source: {error}") from error
    errors = validate_catalog(catalog)
    if errors:
        raise ValueError("; ".join(errors))
    return catalog


FIXED_TIME = datetime(2026, 9, 24)
TEMPLATE_NAMES = (
    "delivery-pack.zh.md", "delivery-pack.en.md",
    "report.zh.docx", "report.en.docx",
    "register.zh.xlsx", "register.en.xlsx",
)


def render_markdown(slug: str, spec: dict, common_fields: list[dict], lang: str) -> str:
    zh = lang == "zh"
    lines = [f"# {spec['title'][lang]}", "", f"Skill: `{slug}`", ""]
    lines.extend([
        f"[{'Word 报告' if zh else 'Word report'}](report.{lang}.docx) · "
        f"[{'Excel 台账' if zh else 'Excel register'}](register.{lang}.xlsx)",
        "", "## " + ("基本信息" if zh else "Record details"), "",
        "| ID | " + ("字段" if zh else "Field") + " | " + ("填写内容" if zh else "Project value") + " |",
        "| --- | --- | --- |",
    ])
    lines.extend(f"| `{field['id']}` | {field[lang]} | |" for field in common_fields)
    lines.extend(["", "## " + ("专题分析与决策" if zh else "Domain analysis and decisions")])
    for section in spec["sections"]:
        lines.extend(["", f"### {section[lang]} (`{section['id']}`)", "", section["instruction"][lang], "", "> " + ("填写项目证据、结论与责任人。" if zh else "Enter project evidence, conclusion, and owner."), ""])
    lines.extend(["", "## " + ("工作台账" if zh else "Working registers")])
    for sheet in spec["sheets"]:
        columns = sheet["columns"]
        lines.extend(["", f"### {sheet['name'][lang]} (`{sheet['id']}`)", "", sheet["purpose"][lang], "", "| " + " | ".join(f"{col[lang]} (`{col['id']}`)" for col in columns) + " |", "| " + " | ".join("---" for _ in columns) + " |", "| " + " | ".join(" " for _ in columns) + " |"])
    return "\n".join(lines) + "\n"


def normalize_ooxml(path: Path) -> None:
    with ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    if path.suffix == ".xlsx":
        for rel_path in sorted(name for name in entries if name.endswith(".rels")):
            relationships = ElementTree.fromstring(entries[rel_path])
            ordered = sorted(relationships, key=lambda relation: (relation.get("Type", ""), relation.get("Target", ""), relation.get("TargetMode", "")))
            replacements = {relation.get("Id"): f"rId{index}" for index, relation in enumerate(ordered, 1)}
            if not replacements:
                continue
            owner = None
            if rel_path != "_rels/.rels":
                prefix, filename = rel_path.rsplit("/_rels/", 1)
                owner = f"{prefix}/{filename[:-5]}"
            for entry_name in (rel_path, owner):
                if entry_name and entry_name in entries:
                    content = entries[entry_name]
                    for old_id, new_id in replacements.items():
                        content = content.replace(f'"{old_id}"'.encode(), f'"{new_id}"'.encode())
                    entries[entry_name] = content
    buffer = BytesIO()
    with ZipFile(buffer, "w") as target:
        for name, contents in sorted(entries.items()):
            info = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, contents)
    path.write_bytes(buffer.getvalue())


def render_docx(slug: str, spec: dict, common_fields: list[dict], lang: str, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    font_name = "Arial Unicode MS" if lang == "zh" else "Arial"
    normal.font.name = font_name
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = font_name
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
        style.font.color.rgb = RGBColor(23, 23, 23)
    title_style = document.styles["Title"]
    title_border = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        title_style.element.get_or_add_pPr().remove(title_border)
    document.core_properties.title = spec["report_title"][lang]
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.add_paragraph(spec["report_title"][lang], "Title")
    subtitle = document.add_paragraph(f"DAMA SKILL · {slug}")
    subtitle.style = "Subtitle"
    document.add_heading("基本信息" if lang == "zh" else "Record details", level=1)
    metadata = document.add_table(rows=0, cols=2)
    metadata.style = "Light Shading Accent 1"
    for field in common_fields:
        cells = metadata.add_row().cells
        cells[0].text = f"{field[lang]} · {field['id']}"
        cells[1].text = " "
    document.add_heading("专题分析与决策" if lang == "zh" else "Domain analysis and decisions", level=1)
    for section_spec in spec["sections"]:
        document.add_heading(f"{section_spec[lang]} · {section_spec['id']}", level=2)
        document.add_paragraph(section_spec["instruction"][lang])
        blank = document.add_paragraph(" ")
        blank.paragraph_format.space_after = Pt(12)
    footer = section.footer.paragraphs[0]
    footer.text = "DAMA · " + ("项目交付模板" if lang == "zh" else "Project deliverable template")
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if lang == "zh":
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
        paragraphs = list(document.paragraphs)
        paragraphs.extend(paragraph for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for paragraph in paragraphs:
            for run in paragraph.runs:
                run.font.name = font_name
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
    document.save(path)
    normalize_ooxml(path)


def build_all(catalog_path: Path, skills_root: Path) -> list[Path]:
    catalog = load_catalog(catalog_path)
    for slug in catalog["skills"]:
        if not (skills_root / slug).is_dir():
            raise ValueError(f"missing skill directory: {slug}")
    outputs: list[Path] = []
    for slug, spec in sorted(catalog["skills"].items()):
        output_dir = skills_root / slug / "templates"
        output_dir.mkdir(exist_ok=True)
        for lang in ("zh", "en"):
            markdown = output_dir / f"delivery-pack.{lang}.md"
            report = output_dir / f"report.{lang}.docx"
            register = output_dir / f"register.{lang}.xlsx"
            markdown.write_text(render_markdown(slug, spec, catalog["common_fields"], lang), encoding="utf-8", newline="\n")
            render_docx(slug, spec, catalog["common_fields"], lang, report)
            outputs.extend((markdown, report, register))
    with tempfile.TemporaryDirectory() as temporary:
        config_path = Path(temporary) / "xlsx-input.json"
        config_path.write_text(json.dumps({"catalog": catalog, "skills_root": str(skills_root)}, ensure_ascii=False), encoding="utf-8")
        node = os.environ.get("DAMA_NODE", "node")
        subprocess.run([node, str(Path(__file__).with_name("build_template_workbooks.mjs")), str(config_path)], check=True)
    for path in outputs:
        if path.suffix == ".xlsx":
            normalize_ooxml(path)
    return outputs


def main() -> int:
    import argparse

    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Build 17 bilingual DAMA skill template sets")
    parser.add_argument("--check", action="store_true", help="Compare generated files with committed templates")
    args = parser.parse_args()
    catalog_path = root / "templates" / "catalog.json"
    skills_root = root / "skills"
    if not args.check:
        print(f"Generated {len(build_all(catalog_path, skills_root))} template files")
        return 0
    with tempfile.TemporaryDirectory() as temporary:
        temporary_root = Path(temporary)
        for slug in EXPECTED_SKILLS:
            (temporary_root / slug).mkdir()
        generated = build_all(catalog_path, temporary_root)
        changed = [str(path.relative_to(temporary_root)) for path in generated if not (skills_root / path.relative_to(temporary_root)).is_file() or path.read_bytes() != (skills_root / path.relative_to(temporary_root)).read_bytes()]
    if changed:
        print("Out-of-date templates: " + ", ".join(changed))
        return 1
    print("OK: 102 generated template files match catalog")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
